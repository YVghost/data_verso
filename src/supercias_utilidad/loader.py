"""
ETL loader: Ranking de Empresas Supercias — Ecuador

Fuente: ranking_{YYYY}.xlsx  (una sola hoja por archivo)

Tabla destino: supercias_utilidad

Estructura del Excel (consistente 2010-presente):
  Fila 1  : metadato "Fecha de corte..."
  Fila 2  : encabezados (22 columnas)
  Fila 3+ : datos

Columnas del Excel → campo BD:
  1  Posición           → posicion       (0 = no rankeado este año)
  2  Año                → anio_ranking   (0 = no rankeado este año)
  3  Posición (ant)     → posicion_ant
  4  Año (ant)          → anio_ranking_ant
  5  Expediente         → expediente
  6  Nombre             → nombre
  7  Tipo Compañia      → tipo_compania
  8  Actividad económica → actividad_eco
  9  Región             → region
  10 Provincia          → provincia
  11 Ciudad             → ciudad
  12 Tamaño             → tamano
  13 Sector             → sector
  14 Cant. Empleados    → cant_empleados
  15 Activo\n{YYYY}     → activo
  16 Patrimonio\n{YYYY} → patrimonio
  17 Ingreso por ventas\n{YYYY} → ingreso_ventas
  18 Utilidad antes del impuesto\n{YYYY} → utilidad_ai
  19 Utilidad del ejercicio\n{YYYY}      → utilidad_ej
  20 Utilidad neta\n{YYYY}              → utilidad_neta
  21 IR causado\n{YYYY}                 → ir_causado
  22 Ingreso Total\n{YYYY}              → ingreso_total

Deduplicación: SHA-256 sobre (anio_archivo, expediente).
"""

import hashlib
import sys
from pathlib import Path

import openpyxl
from sqlalchemy import text

sys.path.append(str(Path(__file__).resolve().parents[2]))
from utils.base_engine import get_master_engine

_TABLE = "supercias_utilidad"

_DDL = """
CREATE TABLE supercias_utilidad (
    id                BIGINT IDENTITY(1,1) NOT NULL,
    anio_archivo      SMALLINT       NOT NULL,
    posicion          INT            NULL,
    anio_ranking      SMALLINT       NULL,
    posicion_ant      INT            NULL,
    anio_ranking_ant  SMALLINT       NULL,
    expediente        INT            NOT NULL,
    nombre            NVARCHAR(500)  NULL,
    tipo_compania     NVARCHAR(200)  NULL,
    actividad_eco     NVARCHAR(MAX)  NULL,
    region            NVARCHAR(50)   NULL,
    provincia         NVARCHAR(100)  NULL,
    ciudad            NVARCHAR(100)  NULL,
    tamano            NVARCHAR(50)   NULL,
    sector            NVARCHAR(100)  NULL,
    cant_empleados    INT            NULL,
    activo            FLOAT          NULL,
    patrimonio        FLOAT          NULL,
    ingreso_ventas    FLOAT          NULL,
    utilidad_ai       FLOAT          NULL,
    utilidad_ej       FLOAT          NULL,
    utilidad_neta     FLOAT          NULL,
    ir_causado        FLOAT          NULL,
    ingreso_total     FLOAT          NULL,
    hash_registro     NVARCHAR(64)   NOT NULL,
    fecha_carga       DATETIME2      DEFAULT GETDATE(),
    CONSTRAINT PK_supercias_utilidad PRIMARY KEY NONCLUSTERED (id)
)"""

_IDX = (
    "CREATE CLUSTERED INDEX CIX_supercias_utilidad "
    "ON supercias_utilidad (anio_archivo, expediente)"
)

_BATCH   = 5_000
_HDR_ROW = 2   # fila de encabezados (1-indexed)
_DAT_ROW = 3   # primera fila de datos


# ---------------------------------------------------------------------------
# Punto de entrada
# ---------------------------------------------------------------------------

def load(files: list[Path]) -> None:
    if not files:
        print("[supercias] Sin archivos para cargar.")
        return

    engine = get_master_engine()
    _ensure_table(engine)
    existing = _get_existing_hashes(engine)

    total = 0
    for path in sorted(files):
        n = _load_file(path, engine, existing)
        total += n

    print(f"[supercias] Total insertado: {total:,} filas en {_TABLE}.")


# ---------------------------------------------------------------------------
# DDL
# ---------------------------------------------------------------------------

def _ensure_table(engine) -> None:
    from sqlalchemy import inspect as sa_inspect
    insp = sa_inspect(engine)
    if not insp.has_table(_TABLE):
        with engine.begin() as conn:
            conn.execute(text(_DDL))
            conn.execute(text(_IDX))
        print(f"[supercias] Tabla {_TABLE} creada.")
    else:
        # Ampliar columnas que pudieron haberse creado con tamaño insuficiente
        cols = {c["name"]: str(c.get("type", "")).upper()
                for c in insp.get_columns(_TABLE)}
        alteraciones = []
        if "MAX" not in cols.get("actividad_eco", ""):
            alteraciones.append(
                f"ALTER TABLE {_TABLE} ALTER COLUMN actividad_eco NVARCHAR(MAX) NULL"
            )
        if cols.get("nombre", "") not in ("NVARCHAR(MAX)", "NVARCHAR(500)"):
            alteraciones.append(
                f"ALTER TABLE {_TABLE} ALTER COLUMN nombre NVARCHAR(500) NULL"
            )
        if cols.get("tipo_compania", "") not in ("NVARCHAR(MAX)", "NVARCHAR(200)"):
            alteraciones.append(
                f"ALTER TABLE {_TABLE} ALTER COLUMN tipo_compania NVARCHAR(200) NULL"
            )
        for ddl_alt in alteraciones:
            with engine.begin() as conn:
                conn.execute(text(ddl_alt))
            col = ddl_alt.split("COLUMN")[1].strip().split()[0]
            print(f"[supercias] Columna {col} ampliada.")


def _get_existing_hashes(engine) -> set[str]:
    with engine.connect() as conn:
        rows = conn.execute(
            text(f"SELECT hash_registro FROM {_TABLE}")
        ).fetchall()
    return {r[0] for r in rows}


# ---------------------------------------------------------------------------
# Por archivo
# ---------------------------------------------------------------------------

def _load_file(path: Path, engine, existing: set[str]) -> int:
    # Extraer año del nombre del archivo (ranking_2025.xlsx → 2025)
    try:
        anio_archivo = int(path.stem.split("_")[-1])
    except ValueError:
        print(f"[supercias] No se pudo extraer año de '{path.name}', omitiendo.")
        return 0

    print(f"[supercias] Procesando {path.name}  (año {anio_archivo})...")
    try:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    except Exception as ex:
        print(f"[supercias] Error abriendo {path.name}: {ex}")
        return 0

    ws = wb.active
    records: list[dict] = []
    inserted = 0

    for row in ws.iter_rows(min_row=_DAT_ROW, values_only=True):
        rec = _parse_row(row, anio_archivo)
        if rec is None:
            continue
        h = _hash(anio_archivo, rec["expediente"])
        if h in existing:
            continue
        rec["hash_registro"] = h
        records.append(rec)
        existing.add(h)

        if len(records) >= _BATCH:
            _insert(records, engine)
            inserted += len(records)
            records = []

    if records:
        _insert(records, engine)
        inserted += len(records)

    wb.close()
    print(f"[supercias] {path.name}: {inserted:,} filas nuevas.")
    return inserted


# ---------------------------------------------------------------------------
# Parser de fila
# ---------------------------------------------------------------------------

def _parse_row(row: tuple, anio_archivo: int) -> dict | None:
    if len(row) < 22:
        return None

    expediente = _to_int(row[4])
    if expediente is None:
        return None

    return {
        "anio_archivo":      anio_archivo,
        "posicion":          _to_int(row[0]),
        "anio_ranking":      _to_int(row[1]),
        "posicion_ant":      _to_int(row[2]),
        "anio_ranking_ant":  _to_int(row[3]),
        "expediente":        expediente,
        "nombre":            _clean(row[5]),
        "tipo_compania":     _clean(row[6]),
        "actividad_eco":     _clean(row[7]),
        "region":            _clean(row[8]),
        "provincia":         _clean(row[9]),
        "ciudad":            _clean(row[10]),
        "tamano":            _clean(row[11]),
        "sector":            _clean(row[12]),
        "cant_empleados":    _to_int(row[13]),
        "activo":            _to_float(row[14]),
        "patrimonio":        _to_float(row[15]),
        "ingreso_ventas":    _to_float(row[16]),
        "utilidad_ai":       _to_float(row[17]),
        "utilidad_ej":       _to_float(row[18]),
        "utilidad_neta":     _to_float(row[19]),
        "ir_causado":        _to_float(row[20]),
        "ingreso_total":     _to_float(row[21]),
    }


# ---------------------------------------------------------------------------
# Inserción
# ---------------------------------------------------------------------------

def _insert(records: list[dict], engine) -> None:
    if not records:
        return
    cols         = list(records[0].keys())
    col_list     = ", ".join(cols)
    placeholders = ", ".join(f":{c}" for c in cols)
    with engine.begin() as conn:
        conn.execute(
            text(f"INSERT INTO {_TABLE} ({col_list}) VALUES ({placeholders})"),
            records,
        )


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _clean(v) -> str | None:
    if v is None:
        return None
    s = str(v).strip()
    return s if s and s.lower() not in ("none", "nan") else None


def _to_int(v) -> int | None:
    if v is None:
        return None
    try:
        n = int(float(str(v).strip()))
        return None if n == 0 and str(v).strip() == "0.0" else n
    except (ValueError, TypeError):
        return None


def _to_float(v) -> float | None:
    if v is None:
        return None
    try:
        import math
        f = float(str(v).strip())
        return None if math.isnan(f) or math.isinf(f) else f
    except (ValueError, TypeError):
        return None


def _hash(anio_archivo: int, expediente: int) -> str:
    key = f"{anio_archivo}|{expediente}"
    return hashlib.sha256(key.encode()).hexdigest()
